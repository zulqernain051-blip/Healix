import json
import os
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

COLOR_NAVY = RGBColor(11, 66, 104)
COLOR_TEXT = RGBColor(30, 41, 59)
COLOR_MUTED = RGBColor(100, 116, 139)

def get_module_sort_order(module_name):
    order = {
        'Common': 2,
        'Auth': 3,
        'Patient': 4,
        'Nurse': 5,
        'Doctor': 6,
        'Admin': 7,
        'Paramedic': 8,
        'Web': 9
    }
    return order.get(module_name, 10)

def set_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "Segoe UI"
        run.font.color.rgb = COLOR_NAVY
    return h

def generate_docx(screens, out_path):
    doc = docx.Document()
    
    # Title Page
    doc.add_paragraph().paragraph_format.space_before = Pt(72)
    p_title = doc.add_paragraph("HEALIX")
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.runs[0]
    r_title.font.name = "Segoe UI"
    r_title.font.size = Pt(48)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_NAVY

    p_sub = doc.add_paragraph("COMPLETE VISUAL USER MANUAL")
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.runs[0].font.name = "Segoe UI"
    p_sub.runs[0].font.size = Pt(16)
    
    p_desc = doc.add_paragraph("Final Year Project Screen-by-Screen Documentation")
    p_desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_desc.runs[0].font.color.rgb = COLOR_MUTED
    doc.add_page_break()

    # Intro
    set_heading(doc, "Chapter 1: Introduction", 1)
    doc.add_paragraph("This document serves as the exhaustive User Manual for the Healix ecosystem. It documents every user-facing screen, derived directly from the source code, outlining what the user sees, how to interact with the interface, and the navigation pathways available.")
    doc.add_page_break()

    # Modules
    modules = {}
    for s in screens:
        m = s.get('module', 'Common')
        if m not in modules: modules[m] = []
        modules[m].append(s)
        
    sorted_modules = sorted(modules.keys(), key=get_module_sort_order)
    screen_index = []
    
    for mod_idx, mod_name in enumerate(sorted_modules, start=2):
        set_heading(doc, f"Chapter {mod_idx}: {mod_name} Module", 1)
        mod_screens = sorted(modules[mod_name], key=lambda x: x['route'].count('/'))
        
        for s_idx, s in enumerate(mod_screens, start=1):
            s_num = f"{mod_idx}.{s_idx}"
            s_title = s['name'].replace('Mobile ', '').replace('Web ', '').replace('/', ' ').title().strip()
            if not s_title: s_title = "Root Layout"
            
            set_heading(doc, f"{s_num} {s_title}", 2)
            
            # Metadata
            p = doc.add_paragraph()
            p.add_run("Screen ID: ").bold = True
            p.add_run(f"{s['id']} | ")
            p.add_run("Platform: ").bold = True
            p.add_run(f"{s['platform'].title()} | ")
            p.add_run("Route: ").bold = True
            p.add_run(f"{s['route']}")
            
            # Screenshot Note
            set_heading(doc, "Screenshot", 3)
            p_ss = doc.add_paragraph("Actual screenshot unavailable — reason: Automated E2E capture not configured for authenticated deep-link states. ")
            p_ss.add_run(f"(Sourced from: {s['sourceFile']})").font.color.rgb = COLOR_MUTED
            
            # Purpose
            set_heading(doc, "Purpose", 3)
            doc.add_paragraph(f"Provides the user interface for the {s['route']} path within the {mod_name} module.")
            
            # What User Sees
            set_heading(doc, "What the User Sees", 3)
            for elem in s.get('ui_elements', []):
                doc.add_paragraph(f"• {elem['name']}: {elem['description']}")
                
            # How to Use
            set_heading(doc, "How to Use the Screen", 3)
            doc.add_paragraph("1. Navigate to this screen.\n2. Review information.\n3. Interact with controls.\n4. System processes actions.")
            
            # Action -> Result
            set_heading(doc, "Action → Result", 3)
            if s.get('ui_actions'):
                t = doc.add_table(rows=1, cols=2)
                t.style = 'Table Grid'
                t.rows[0].cells[0].text = "User Action"
                t.rows[0].cells[1].text = "Result"
                for action in s['ui_actions']:
                    row = t.add_row()
                    row.cells[0].text = action['action']
                    row.cells[1].text = action['result']
            else:
                doc.add_paragraph("No specific interactive actions detected.")
                
            # Navigation
            set_heading(doc, "Navigation", 3)
            doc.add_paragraph(f"Previous Screen ↓\nCURRENT SCREEN ({s_title}) ↓")
            for action in s.get('ui_actions', []):
                if 'Navigates' in action['result']:
                    doc.add_paragraph(f"• Next: {action['result'].replace('Navigates to ', '')}")
                    
            doc.add_paragraph("_" * 40)
            screen_index.append([s_num, s_title, mod_name, mod_name, s['platform'].title(), s['route']])
            
        doc.add_page_break()

    # Screen Index
    set_heading(doc, "Chapter 10: Screen Index", 1)
    t_idx = doc.add_table(rows=1, cols=6)
    t_idx.style = 'Table Grid'
    headers = ["No.", "Name", "Module", "Role", "Platform", "Route"]
    for i, h in enumerate(headers):
        t_idx.rows[0].cells[i].text = h
    for row_data in screen_index:
        row = t_idx.add_row()
        for i, val in enumerate(row_data):
            row.cells[i].text = str(val)
            
    doc.save(out_path)
    print(f"Generated User Manual DOCX at {out_path}")

def run():
    base_dir = r"F:\class Data\FYP Project\Proposal\Project\Healix"
    inv_path = os.path.join(base_dir, "documentation", "inventory", "screens.json")
    out_docx = os.path.join(base_dir, "documentation", "output", "Healix_User_Manual.docx")
    
    with open(inv_path, 'r', encoding='utf-8') as f:
        screens = json.load(f)
        
    generate_docx(screens, out_docx)
    
if __name__ == "__main__":
    run()
